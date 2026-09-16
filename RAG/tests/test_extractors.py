import io
import zipfile

import pytest
from docx import Document as DocxDocument

from app.services import extractors
from app.exceptions.custom_exceptions import (
    DocumentExtractionException,
    InvalidZipException,
    UnsupportedFileTypeException,
)


def test_extract_txt():
    content = "Hello, this is a plain text file.".encode("utf-8")
    assert extractors.extract_text("notes.txt", content) == "Hello, this is a plain text file."


def test_extract_docx_reads_real_paragraphs():
    document = DocxDocument()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")

    buffer = io.BytesIO()
    document.save(buffer)

    text = extractors.extract_text("notes.docx", buffer.getvalue())

    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_extract_corrupted_pdf_raises_clear_error():
    with pytest.raises(DocumentExtractionException):
        extractors.extract_text("broken.pdf", b"this is not a real pdf file")


def test_extract_unsupported_extension_raises():
    with pytest.raises(UnsupportedFileTypeException):
        extractors.extract_text("script.exe", b"MZ")


def _make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        for name, content in entries.items():
            zf.writestr(name, content)
    return buffer.getvalue()


def test_zip_extraction_returns_supported_files():
    zip_bytes = _make_zip({"a.txt": b"hello", "b.pdf": b"not a real pdf but that's ok here"})

    files = extractors.extract_files_from_zip(zip_bytes)

    names = {name for name, _ in files}
    assert names == {"a.txt", "b.pdf"}


def test_zip_extraction_flattens_path_traversal_attempts():
    """A malicious filename like '../../evil.txt' must come back as just
    'evil.txt' - we never trust the archive's own directory structure."""
    zip_bytes = _make_zip({"../../evil.txt": b"payload"})

    files = extractors.extract_files_from_zip(zip_bytes)

    assert files[0][0] == "evil.txt"
    assert "/" not in files[0][0] and ".." not in files[0][0]


def test_zip_extraction_skips_unsupported_files():
    zip_bytes = _make_zip({"a.txt": b"hello", "virus.exe": b"MZ"})

    files = extractors.extract_files_from_zip(zip_bytes)

    names = {name for name, _ in files}
    assert names == {"a.txt"}


def test_zip_with_only_unsupported_files_raises():
    zip_bytes = _make_zip({"virus.exe": b"MZ"})

    with pytest.raises(InvalidZipException):
        extractors.extract_files_from_zip(zip_bytes)


def test_empty_zip_raises():
    zip_bytes = _make_zip({})

    with pytest.raises(InvalidZipException):
        extractors.extract_files_from_zip(zip_bytes)


def test_corrupted_zip_raises():
    with pytest.raises(InvalidZipException):
        extractors.extract_files_from_zip(b"not a zip file at all")
